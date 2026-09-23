import json
import os
import sys

from text_utils import dividir_em_chunks, limpar_texto

OUTROS_DIR = "outros"
URLS_FILE = "urls.txt"
WIKI_FILE = "wikipedia.txt"
WIKI_LANG_PADRAO = "pt"
OUT_FILE = "data/documents_outros.jsonl"

LINHAS_POR_BLOCO_PLANILHA = 20


def extrair_txt_md(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        texto = f.read()
    return limpar_texto(texto)


def extrair_docx(path: str) -> str:
    try:
        import docx
    except ImportError:
        print(
            "  [aviso] biblioteca 'python-docx' não instalada; pulando "
            f"'{path}'. Rode: pip install python-docx"
        )
        return ""

    try:
        documento = docx.Document(path)
    except Exception as e:
        print(f"  [aviso] não consegui abrir '{path}': {e}")
        return ""

    texto = "\n".join(p.text for p in documento.paragraphs)
    return limpar_texto(texto)


def extrair_planilha(path: str) -> list[str]:
    ext = os.path.splitext(path)[1].lower()
    linhas_texto: list[str] = []

    try:
        if ext == ".csv":
            import csv

            with open(path, "r", encoding="utf-8", errors="ignore", newline="") as f:
                leitor = csv.reader(f)
                cabecalho = next(leitor, None)
                for linha in leitor:
                    if cabecalho:
                        pares = [f"{c}: {v}" for c, v in zip(cabecalho, linha)]
                    else:
                        pares = linha
                    if pares:
                        linhas_texto.append(" | ".join(pares))
        else:
            try:
                from openpyxl import load_workbook
            except ImportError:
                print(
                    "  [aviso] biblioteca 'openpyxl' não instalada; pulando "
                    f"'{path}'. Rode: pip install openpyxl"
                )
                return []

            wb = load_workbook(path, read_only=True, data_only=True)
            for aba in wb.worksheets:
                linhas = list(aba.iter_rows(values_only=True))
                if not linhas:
                    continue
                cabecalho = linhas[0]
                for linha in linhas[1:]:
                    pares = [
                        f"{c}: {v}" for c, v in zip(cabecalho, linha) if v is not None
                    ]
                    if pares:
                        linhas_texto.append(" | ".join(pares))
    except Exception as e:
        print(f"  [aviso] erro lendo planilha '{path}': {e}")
        return []

    return linhas_texto


def extrair_srt(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        linhas = f.readlines()

    falas = []
    for linha in linhas:
        linha = linha.strip()
        if not linha:
            continue
        if linha.isdigit():
            continue
        if "-->" in linha:
            continue
        falas.append(linha)

    return limpar_texto(" ".join(falas))


def buscar_titulos_wikipedia(
    termo: str, lang: str = WIKI_LANG_PADRAO, limite: int = 5
) -> list[str]:
    try:
        import requests
    except ImportError:
        print(
            "  [aviso] biblioteca 'requests' não instalada; pulando busca na Wikipédia."
        )
        return []

    try:
        resp = requests.get(
            f"https://{lang}.wikipedia.org/w/api.php",
            params={
                "action": "query",
                "list": "search",
                "srsearch": termo,
                "format": "json",
                "srlimit": limite,
            },
            timeout=10,
            headers={"User-Agent": "Celina/1.0"},
        )
        resp.raise_for_status()
        dados = resp.json()
        return [item["title"] for item in dados.get("query", {}).get("search", [])]
    except Exception as e:
        print(f"  [aviso] busca na Wikipédia falhou: {e}")
        return []


def extrair_wikipedia(titulo: str, lang: str = WIKI_LANG_PADRAO) -> str:
    try:
        import requests
    except ImportError:
        print("  [aviso] biblioteca 'requests' não instalada; pulando Wikipédia.")
        return ""

    try:
        resp = requests.get(
            f"https://{lang}.wikipedia.org/w/api.php",
            params={
                "action": "query",
                "prop": "extracts",
                "explaintext": 1,
                "titles": titulo,
                "format": "json",
                "redirects": 1,
            },
            timeout=15,
            headers={"User-Agent": "Celina/1.0"},
        )
        resp.raise_for_status()
        paginas = resp.json().get("query", {}).get("pages", {})
        for pagina in paginas.values():
            texto = pagina.get("extract", "")
            if texto:
                return limpar_texto(texto)
    except Exception as e:
        print(f"  [aviso] não consegui baixar o artigo '{titulo}' da Wikipédia: {e}")

    return ""


def extrair_pagina_web(url: str) -> str:
    try:
        import requests
        from bs4 import BeautifulSoup
    except ImportError:
        print(
            "  [aviso] bibliotecas 'requests'/'beautifulsoup4' não instaladas; "
            "pulando URLs. Rode: pip install requests beautifulsoup4"
        )
        return ""

    try:
        resp = requests.get(url, timeout=15, headers={"User-Agent": "Mozilla/5.0"})
        resp.raise_for_status()
    except Exception as e:
        print(f"  [aviso] não consegui baixar '{url}': {e}")
        return ""

    soup = BeautifulSoup(resp.text, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header", "noscript"]):
        tag.decompose()

    texto = soup.get_text(separator=" ")
    return limpar_texto(texto)


def montar_chunks_texto(source_name: str, texto: str, resultados: list[dict]) -> None:
    for i, chunk in enumerate(dividir_em_chunks(texto)):
        if len(chunk) < 20:
            continue
        resultados.append({"source": source_name, "page": 1, "chunk": i, "text": chunk})


def montar_chunks_planilha(
    source_name: str, linhas: list[str], resultados: list[dict]
) -> None:
    bloco: list[str] = []
    bloco_num = 0
    for linha in linhas:
        bloco.append(linha)
        if len(bloco) >= LINHAS_POR_BLOCO_PLANILHA:
            bloco_num += 1
            resultados.append(
                {
                    "source": source_name,
                    "page": bloco_num,
                    "chunk": 0,
                    "text": "\n".join(bloco),
                }
            )
            bloco = []
    if bloco:
        bloco_num += 1
        resultados.append(
            {
                "source": source_name,
                "page": bloco_num,
                "chunk": 0,
                "text": "\n".join(bloco),
            }
        )


def processar_pasta_outros(resultados: list[dict]) -> None:
    if not os.path.isdir(OUTROS_DIR):
        return

    for fname in sorted(os.listdir(OUTROS_DIR)):
        path = os.path.join(OUTROS_DIR, fname)
        if not os.path.isfile(path):
            continue
        ext = os.path.splitext(fname)[1].lower()

        if ext in (".txt", ".md"):
            texto = extrair_txt_md(path)
            if texto:
                montar_chunks_texto(fname, texto, resultados)
        elif ext == ".srt":
            texto = extrair_srt(path)
            if texto:
                montar_chunks_texto(fname, texto, resultados)
        elif ext == ".docx":
            texto = extrair_docx(path)
            if texto:
                montar_chunks_texto(fname, texto, resultados)
        elif ext in (".csv", ".xlsx"):
            linhas = extrair_planilha(path)
            if linhas:
                montar_chunks_planilha(fname, linhas, resultados)
        else:
            print(f"  [aviso] formato '{ext}' não suportado, ignorando '{fname}'.")


def processar_urls(resultados: list[dict]) -> None:
    if not os.path.exists(URLS_FILE):
        return

    with open(URLS_FILE, "r", encoding="utf-8") as f:
        urls = [
            linha.strip()
            for linha in f
            if linha.strip() and not linha.strip().startswith("#")
        ]

    for url in urls:
        texto = extrair_pagina_web(url)
        if texto:
            montar_chunks_texto(url, texto, resultados)


def processar_wikipedia(resultados: list[dict]) -> None:
    if not os.path.exists(WIKI_FILE):
        return

    with open(WIKI_FILE, "r", encoding="utf-8") as f:
        titulos = [
            linha.strip()
            for linha in f
            if linha.strip() and not linha.strip().startswith("#")
        ]

    for titulo in titulos:
        texto = extrair_wikipedia(titulo)
        if texto:
            montar_chunks_texto(f"Wikipédia: {titulo}", texto, resultados)


def main() -> int:
    os.makedirs("data", exist_ok=True)

    resultados: list[dict] = []
    processar_pasta_outros(resultados)
    processar_urls(resultados)
    processar_wikipedia(resultados)

    if not resultados:
        print(
            f"Nenhum conteúdo encontrado em '{OUTROS_DIR}/', '{URLS_FILE}' nem '{WIKI_FILE}'."
        )
        return 1

    with open(OUT_FILE, "w", encoding="utf-8") as f:
        for d in resultados:
            f.write(json.dumps(d, ensure_ascii=False) + "\n")

    fontes = {d["source"] for d in resultados}
    print(
        f"Extraídos {len(resultados)} blocos de {len(fontes)} fonte(s) para {OUT_FILE}"
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
